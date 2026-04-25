#!/usr/bin/env python3
"""Inbox scanner — extract text, classify signals, generate metadata."""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

# ── paths ──────────────────────────────────────────────────────────────

INBOX_ROOT = Path(__file__).resolve().parent.parent  # inbox/
RAW_DIR = INBOX_ROOT / "raw"
PROCESSED_DIR = INBOX_ROOT / "processed"
INDEX_PATH = PROCESSED_DIR / "index.json"

CATEGORIES = [
    "regulatory", "industry", "products", "markets",
    "cases", "professional", "tech-digital",
]

# ── classification rules (mirrors build-context.py patterns) ───────────

@dataclass(frozen=True)
class SignalRule:
    pattern: str
    signal_type: str
    label_zh: str


SIGNAL_RULES: list[SignalRule] = [
    SignalRule(r"政策|监管|通知|办法|意见|规则|条例|央行|金融监管|证监|交易商协会|银保监|人民银行", "policy", "政策监管"),
    SignalRule(r"新闻|报道|发布|公告|采访|会议|论坛|发布会|白皮书", "news", "新闻动态"),
    SignalRule(r"产品|方案|服务|功能|办理|申请|额度|融资|结算|账户|手续费|费率", "product", "产品方案"),
    SignalRule(r"竞品|同业|他行|招行|工行|建行|农行|中行|平安|微众|网商", "competitor", "同业竞品"),
    SignalRule(r"客户|企业|公司|拜访|订单|项目|扩产|减产|股权|中标|诉讼", "client", "客户线索"),
    SignalRule(r"利率|汇率|通胀|经济|行业|周期|价格|大宗|出口|内需|GDP|CPI|PMI", "market", "市场环境"),
    SignalRule(r"风险|合规|反洗钱|KYC|逾期|不良|诉讼|处罚|罚单", "risk", "风险合规"),
    SignalRule(r"方法|技能|复盘|经验|话术|表达|框架|学习|认证|管理|沟通", "capability", "能力素材"),
]


def classify_signal(text: str) -> tuple[str | None, list[str]]:
    """Return (primary_signal_type, all_matching_types)."""
    matches: list[str] = []
    for rule in SIGNAL_RULES:
        if re.search(rule.pattern, text):
            matches.append(rule.signal_type)
    primary = matches[0] if matches else None
    return primary, matches


# ── text extraction ────────────────────────────────────────────────────

def read_text_file(path: Path) -> str:
    """Read .md / .txt files."""
    return path.read_text(encoding="utf-8", errors="replace")


def read_url_file(path: Path) -> tuple[str | None, str]:
    """Read .url.txt — return (url, fetched_text)."""
    url = path.read_text(encoding="utf-8").strip().splitlines()[0].strip()
    try:
        import trafilatura
        downloaded = trafilatura.fetch_url(url)
        if downloaded:
            text = trafilatura.extract(downloaded) or ""
            return url, text
    except ImportError:
        pass
    try:
        import urllib.request
        with urllib.request.urlopen(url, timeout=15) as resp:
            raw = resp.read(200_000).decode("utf-8", errors="replace")
            text = re.sub(r"<[^>]+>", " ", raw)
            text = re.sub(r"\s+", " ", text).strip()
            return url, text[:10_000]
    except Exception:
        return url, ""


def read_pdf(path: Path) -> str:
    """Extract text from PDF (best-effort)."""
    try:
        import PyPDF2
        reader = PyPDF2.PdfReader(str(path))
        pages = [p.extract_text() or "" for p in reader.pages[:30]]
        return "\n".join(pages)
    except ImportError:
        pass
    try:
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            pages = [p.extract_text() or "" for p in pdf.pages[:30]]
            return "\n".join(pages)
    except ImportError:
        pass
    return "[PDF 需要安装 PyPDF2 或 pdfplumber 才能提取文本]"


def read_docx(path: Path) -> str:
    """Extract text from .docx."""
    try:
        import docx
        doc = docx.Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    except ImportError:
        return "[DOCX 需要安装 python-docx 才能提取文本]"


def read_spreadsheet(path: Path) -> str:
    """Extract headers + first rows from .xlsx / .csv."""
    if path.suffix == ".csv":
        try:
            import csv
            with open(path, encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                rows = [next(reader, []) for _ in range(20)]
            return "\n".join(",".join(r) for r in rows)
        except Exception:
            return "[CSV 读取失败]"
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        ws = wb.active
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 20:
                break
            rows.append(",".join(str(c) if c is not None else "" for c in row))
        return "\n".join(rows)
    except ImportError:
        return "[XLSX 需要安装 openpyxl 才能提取文本]"


def extract_text(path: Path) -> tuple[str, str | None]:
    """Extract text from a file. Returns (text, source_url_or_None)."""
    suffix = path.suffix.lower()
    name = path.name.lower()

    if name.endswith(".url.txt"):
        url, text = read_url_file(path)
        return text, url
    if suffix in (".md", ".txt", ".markdown"):
        return read_text_file(path), None
    if suffix == ".pdf":
        return read_pdf(path), None
    if suffix == ".docx":
        return read_docx(path), None
    if suffix in (".xlsx", ".csv"):
        return read_spreadsheet(path), None
    try:
        return read_text_file(path), None
    except Exception:
        return "", None


# ── metadata generation ────────────────────────────────────────────────

def next_id(category: str, date_str: str, existing_ids: set[str]) -> str:
    """Generate next sequential ID: inbox-{category}-{YYYYMMDD}-{NNN}."""
    seq = 1
    while True:
        candidate = f"inbox-{category}-{date_str}-{seq:03d}"
        if candidate not in existing_ids:
            return candidate
        seq += 1


def build_meta(
    *,
    entry_id: str,
    filename: str,
    category: str,
    source_type: str,
    source_url: str | None,
    signal_type: str | None,
    text_preview: str,
    now_iso: str,
    source_verified: bool = True,
) -> dict:
    """Build a meta.json dict conforming to inbox-entry.schema.json."""
    return {
        "id": entry_id,
        "filename": filename,
        "category": category,
        "source_type": source_type,
        "source_url": source_url,
        "signal_type": signal_type,
        "text_preview": text_preview[:500],
        "created_at": now_iso,
        "scanned_at": now_iso,
        "status": "scanned",
        "distill_target": None,
        "distill_visibility": None,
        "distill_output_path": None,
        "source_verified": source_verified,
        "notes": None,
    }


# ── index management ──────────────────────────────────────────────────

def load_index() -> dict:
    if INDEX_PATH.exists():
        return json.loads(INDEX_PATH.read_text(encoding="utf-8"))
    return {"version": 1, "last_scan": None, "entries": []}


def save_index(index: dict) -> None:
    INDEX_PATH.parent.mkdir(parents=True, exist_ok=True)
    INDEX_PATH.write_text(json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8")


# ── main scan logic ────────────────────────────────────────────────────

def scan_file(path: Path, category: str, existing_ids: set[str], dry_run: bool) -> dict | None:
    """Scan a single file, return meta dict (or None if skipped)."""
    name = path.name
    if name.startswith(".") or name == ".gitkeep":
        return None
    meta_path = PROCESSED_DIR / f"{name}.meta.json"
    if meta_path.exists():
        print(f"  skip (already scanned): {name}")
        return None

    print(f"  scanning: {name}")
    text, source_url = extract_text(path)
    if not text.strip():
        print(f"    warning: no text extracted from {name}")

    signal_type, _ = classify_signal(text) if text.strip() else (None, [])
    source_type = "url" if source_url else "file"

    date_str = datetime.now().strftime("%Y%m%d")
    entry_id = next_id(category, date_str, existing_ids)
    now_iso = datetime.now(timezone.utc).isoformat()

    meta = build_meta(
        entry_id=entry_id,
        filename=name,
        category=category,
        source_type=source_type,
        source_url=source_url,
        signal_type=signal_type,
        text_preview=text[:500],
        now_iso=now_iso,
        source_verified=bool(text.strip()),
    )

    if not dry_run:
        meta_path.parent.mkdir(parents=True, exist_ok=True)
        meta_path.write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"    -> {meta_path.relative_to(INBOX_ROOT)}")
    else:
        print(f"    -> (dry-run) would write {meta_path.relative_to(INBOX_ROOT)}")

    return meta


def collect_existing_ids() -> set[str]:
    """Collect all existing entry IDs to avoid collisions."""
    ids: set[str] = set()
    for f in PROCESSED_DIR.glob("*.meta.json"):
        try:
            data = json.loads(f.read_text(encoding="utf-8"))
            if "id" in data:
                ids.add(data["id"])
        except Exception:
            pass
    return ids


def scan_category(category: str, dry_run: bool) -> list[dict]:
    """Scan all files in a single category directory."""
    cat_dir = RAW_DIR / category
    if not cat_dir.exists():
        print(f"category '{category}' not found: {cat_dir}")
        return []

    print(f"\n[{category}]")
    files = sorted(f for f in cat_dir.iterdir() if f.is_file())
    if not files:
        print("  (empty)")
        return []

    existing_ids = collect_existing_ids()
    results = []
    for f in files:
        meta = scan_file(f, category, existing_ids, dry_run)
        if meta:
            existing_ids.add(meta["id"])
            results.append(meta)
    return results


def rebuild_index() -> None:
    """Rebuild index.json from all .meta.json files in processed/."""
    index = load_index()
    index["last_scan"] = datetime.now(timezone.utc).isoformat()
    all_entries = []
    for f in sorted(PROCESSED_DIR.glob("*.meta.json")):
        try:
            data = json.loads(f.read_text(encoding="utf-8"))
            all_entries.append({
                "id": data["id"],
                "filename": data["filename"],
                "category": data["category"],
                "status": data["status"],
                "created_at": data["created_at"],
            })
        except Exception:
            pass
    index["entries"] = all_entries
    save_index(index)


def main() -> None:
    parser = argparse.ArgumentParser(description="Inbox scanner — extract text, classify, generate metadata")
    parser.add_argument("--category", "-c", choices=CATEGORIES, help="Scan a specific category")
    parser.add_argument("--all", "-a", action="store_true", help="Scan all categories")
    parser.add_argument("--file", "-f", help="Scan a single file (path relative to inbox/)")
    parser.add_argument("--text", "-t", help="Classify inline text (no file needed)")
    parser.add_argument("--dry-run", action="store_true", help="Preview without writing meta files")
    args = parser.parse_args()

    # inline text classification mode
    if args.text:
        signal_type, all_types = classify_signal(args.text)
        print(json.dumps({
            "signal_type": signal_type,
            "all_matches": all_types,
            "preview": args.text[:200],
        }, ensure_ascii=False, indent=2))
        return

    # file mode
    if args.file:
        fpath = Path(args.file)
        if not fpath.exists():
            fpath = INBOX_ROOT / args.file
        if not fpath.exists():
            print(f"file not found: {args.file}", file=sys.stderr)
            sys.exit(1)
        parts = fpath.parts
        category = "other"
        for part in parts:
            if part in CATEGORIES:
                category = part
                break
        existing_ids = collect_existing_ids()
        scan_file(fpath, category, existing_ids, args.dry_run)
        return

    # category / all mode
    if args.all:
        categories = CATEGORIES
    elif args.category:
        categories = [args.category]
    else:
        parser.print_help()
        sys.exit(1)

    total = 0
    for cat in categories:
        results = scan_category(cat, args.dry_run)
        total += len(results)

    if not args.dry_run and total > 0:
        rebuild_index()

    print(f"\nDone. {total} file(s) scanned.")


if __name__ == "__main__":
    main()
